from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_REPO = ROOT / "docs" / "Traffic_Sign_Legal_QA_Presentation_Script.docx"
OUT_DOWNLOADS = Path("/Users/leloc/Downloads/Traffic_Sign_Legal_QA_Presentation_Script.docx")

BLUE = RGBColor(0x0B, 0x3A, 0x78)
DARK = RGBColor(0x11, 0x1C, 0x2E)
MUTED = RGBColor(0x55, 0x63, 0x75)


SLIDES = [
    {
        "title": "Slide 1 - Project Title",
        "goal": "Giới thiệu ngắn gọn tên đề tài, bài toán và output chính.",
        "script": (
            "Kính chào thầy cô và các bạn. Nhóm em xin trình bày project "
            "Traffic Sign Legal Question Answering. Đây là hệ thống trả lời câu hỏi pháp lý "
            "về biển báo giao thông từ ảnh và câu hỏi tự do bằng tiếng Việt. Input của hệ thống "
            "là một ảnh giao thông và câu hỏi người dùng. Output không chỉ là câu trả lời, mà còn "
            "có citation theo điều luật, evidence pháp lý và evidence thị giác từ biển báo được phát hiện."
        ),
        "transition": "Để thấy vì sao bài toán này cần thiết, em bắt đầu từ motivation của project.",
        "avoid": "Không nói hệ thống thay thế tư vấn pháp lý chính thức; chỉ nói đây là prototype nghiên cứu và hỗ trợ học tập.",
    },
    {
        "title": "Slide 2 - Problem Motivation",
        "goal": "Nêu vấn đề thực tế mà không phóng đại.",
        "script": (
            "Trong thực tế, người tham gia giao thông có thể gặp một biển báo nhưng không chắc biển đó áp dụng cho phương tiện nào, "
            "trong khung giờ nào, hoặc trong trường hợp nào. Nếu chỉ nhìn biển báo thì thiếu căn cứ pháp lý. Nếu chỉ đọc văn bản luật "
            "thì người dùng lại khó map văn bản với tình huống trong ảnh. Vì vậy, project này tập trung vào bài toán kết hợp thị giác, "
            "truy xuất luật và sinh câu trả lời có căn cứ."
        ),
        "transition": "Từ vấn đề đó, slide tiếp theo mô tả một tình huống người dùng cụ thể.",
        "avoid": "Không nói người dùng hoàn toàn không cần đọc luật gốc; nên nói hệ thống giúp định hướng và hiển thị căn cứ để kiểm tra.",
    },
    {
        "title": "Slide 3 - User Scenario",
        "goal": "Giải thích Figure 1 theo luồng người dùng.",
        "script": (
            "Một scenario đơn giản là người dùng chụp hoặc upload ảnh có biển báo giao thông, sau đó hỏi một câu bằng tiếng Việt, "
            "ví dụ xe máy có được đi vào đường này không, hoặc biển này cấm dừng hay cấm đỗ trong trường hợp nào. Hệ thống phân tích ảnh, "
            "phát hiện biển báo, truy xuất điều luật liên quan, rồi trả lời bằng tiếng Việt kèm citation và evidence. Điểm quan trọng ở đây "
            "là người dùng không cần chọn A/B/C/D; câu hỏi và câu trả lời đều ở dạng tự do hơn."
        ),
        "transition": "Sau scenario, em sẽ nói rõ objective và scope để tránh hiểu sai phạm vi của hệ thống.",
        "avoid": "Không hứa hệ thống xử lý mọi loại tình huống giao thông ngoài phạm vi ảnh biển báo và dữ liệu luật đã chuẩn bị.",
    },
    {
        "title": "Slide 4 - Objectives and Scope",
        "goal": "Chốt phạm vi vừa đủ để bảo vệ.",
        "script": (
            "Objective chính của project là xây dựng một pipeline end-to-end cho image-based legal QA về biển báo giao thông. "
            "Trong scope hiện tại, nhóm tập trung vào bốn điểm: phát hiện biển báo trong ảnh, truy xuất các điều luật liên quan, "
            "sinh câu trả lời tiếng Việt dạng free-form, và kiểm soát output bằng JSON có citation, legal evidence và visual evidence. "
            "Nhóm không đặt mục tiêu huấn luyện một mô hình luật từ đầu, mà ưu tiên thiết kế một hệ thống retrieval-grounded, dễ kiểm tra và dễ demo."
        ),
        "transition": "Để pipeline có thể cite được điều luật, bước đầu tiên là chuẩn bị LawDB và dữ liệu biển báo.",
        "avoid": "Không nói model tự học luật; nên nói model được cung cấp context luật và bị ràng buộc chỉ dùng context đó.",
    },
    {
        "title": "Slide 5 - Dataset and LawDB",
        "goal": "Nêu dữ liệu chính và con số an toàn.",
        "script": (
            "Dữ liệu pháp lý của project được tổ chức thành LawDB có cấu trúc. Sau preprocessing, hệ thống có 2 văn bản luật, "
            "402 article-level records và 489 detailed sign entries được chuẩn bị để phục vụ retrieval và hiển thị evidence. "
            "Mỗi article có khóa ổn định gồm law_id và article_id. Cách tổ chức này giúp hệ thống không chỉ lấy text luật, "
            "mà còn kiểm tra citation có thật sự nằm trong corpus hay không."
        ),
        "transition": "Slide tiếp theo trình bày cụ thể hơn cách raw LawDB được biến thành dữ liệu có thể dùng trong hệ thống.",
        "avoid": "Không nói toàn bộ luật giao thông Việt Nam đã được bao phủ; chỉ nói corpus hiện tại được chuẩn bị cho phạm vi project.",
    },
    {
        "title": "Slide 6 - LawDB Preprocessing Workflow",
        "goal": "Giải thích Figure về raw LawDB sang collection.",
        "script": (
            "Ở bước LawDB preprocessing, raw JSON được làm sạch và flatten thành từng article. Với các hình biển báo trong LawDB, "
            "hệ thống tách crop, lưu đường dẫn ảnh, tên biển và mô tả ý nghĩa của biển. Một phần mô tả được rà soát thủ công để giảm lỗi "
            "khi các biển gần giống nhau, ví dụ các biến thể cấm dừng, cấm đỗ, ngày chẵn hoặc ngày lẻ. Sau đó, các thông tin như law_id, "
            "article_id, image_path, sign_name và sign_description được dùng làm metadata cho retrieval và evidence display."
        ),
        "transition": "Sau khi dữ liệu đã sẵn sàng, em chuyển sang pipeline xử lý một câu hỏi mới.",
        "avoid": "Không nói mô tả biển hoàn toàn do model tự suy luận; nên nói có bước chuẩn hóa và review để ổn định chất lượng.",
    },
    {
        "title": "Slide 7 - Multi-stage Pipeline",
        "goal": "Giải thích pipeline tổng thể từ ảnh đến JSON answer.",
        "script": (
            "Pipeline của hệ thống gồm nhiều stage. Đầu tiên, ảnh upload được đưa qua detector để tìm các biển báo và tạo crop. "
            "Sau đó, hệ thống lọc biển liên quan theo câu hỏi, tạo embedding cho câu hỏi và biển báo, rồi truy vấn Qdrant để lấy các item gần nhất. "
            "Các payload trả về cung cấp article_id và metadata biển báo. Tiếp theo, LawDBArticleLookup resolve các ID này thành text điều luật đầy đủ. "
            "Cuối cùng, LLM nhận ảnh, câu hỏi, sign evidence và legal evidence để sinh strict JSON answer."
        ),
        "transition": "Nếu nhìn ở mức kiến trúc phần mềm, các bước này được chia thành bốn layer chính.",
        "avoid": "Không nói retrieval luôn đúng; nên nói retrieval cung cấp candidate evidence và answer prompt yêu cầu model bám vào evidence.",
    },
    {
        "title": "Slide 8 - Final System Architecture",
        "goal": "Giải thích Figure architecture sau khi đã chỉnh mũi tên.",
        "script": (
            "Architecture cuối cùng gồm UI layer, vision layer, retrieval layer và reasoning layer. UI layer là Streamlit, cho phép upload ảnh, nhập câu hỏi "
            "và xem debug panel. Vision layer gồm YOLOE detector và thư mục crop biển báo. Retrieval layer dùng CLIP embedding model và Qdrant collection để lấy metadata liên quan. "
            "Reasoning layer gồm LawDBArticleLookup và một OpenAI-compatible LLM endpoint. Chỉ reasoning layer mới trực tiếp tạo final output; các layer trước cung cấp dữ liệu trung gian cho bước này."
        ),
        "transition": "Điểm khác biệt của hệ thống là output không phải một lựa chọn cố định, mà là một JSON answer có cấu trúc.",
        "avoid": "Khi trình bày figure, tránh nói mọi layer đều trực tiếp trả kết quả; nhấn mạnh dữ liệu chảy tuần tự đến reasoning layer.",
    },
    {
        "title": "Slide 9 - Free-form JSON Answer",
        "goal": "Làm rõ output và guardrail prompt.",
        "script": (
            "Output của hệ thống được chuẩn hóa thành strict JSON gồm answer, citations, visual_evidence, reasoning và confidence. "
            "Trong đó, citations phải có law_id, article_id, title và đoạn evidence pháp lý được dùng để trả lời. visual_evidence ghi lại biển báo được phát hiện hoặc truy xuất, "
            "kèm mô tả và tên ảnh crop. Prompt cũng yêu cầu model chỉ dùng context được cung cấp, không bịa điều luật, và nếu thiếu căn cứ thì phải trả lời là chưa đủ thông tin với confidence thấp."
        ),
        "transition": "Các field này được hiển thị trực tiếp trong Streamlit demo để người xem có thể kiểm tra câu trả lời.",
        "avoid": "Không nói JSON guarantee tuyệt đối từ model; nên nói hệ thống có parser, validation và fallback khi output lỗi.",
    },
    {
        "title": "Slide 10 - Streamlit Demo",
        "goal": "Chuẩn bị lời dẫn demo trực tiếp.",
        "script": (
            "Trong demo, người dùng upload ảnh jpg hoặc png, nhập câu hỏi tự do và bấm Analyze. UI hiển thị ảnh gốc, các crop biển báo, biển nào được xem là relevant, "
            "câu trả lời tiếng Việt, citations và legal evidence trong các expandable sections. Khi bật debug, hệ thống cũng hiển thị raw JSON, retrieved articles, timing và dữ liệu trung gian. "
            "Phần debug giúp nhóm kiểm tra pipeline, còn phần answer/citation là phần dành cho người dùng cuối."
        ),
        "transition": "Để demo chạy được ổn định, nhóm triển khai các service trên máy GPU thuê theo sơ đồ deployment.",
        "avoid": "Trong lúc demo, không nên mở quá nhiều debug field nếu không được hỏi; chỉ mở khi cần chứng minh citation hoặc evidence.",
    },
    {
        "title": "Slide 11 - GPU Deployment",
        "goal": "Giải thích deployment vừa đủ, không sa đà vận hành.",
        "script": (
            "Khi chạy demo online, trình duyệt truy cập public Streamlit URL được map tới port 8501 trên GPU host. Streamlit gọi Python pipeline cục bộ trên cùng máy. "
            "Pipeline gọi Qdrant tại port 6333 để retrieval và gọi Ollama qua OpenAI-compatible endpoint tại port 18989 để sinh câu trả lời. "
            "LLM chạy trên RTX GPU nên có thể xử lý prompt có ảnh, câu hỏi và legal context trong thời gian chấp nhận được cho demo."
        ),
        "transition": "Sau phần hệ thống, em chuyển sang cách nhóm đánh giá chất lượng.",
        "avoid": "Không cần nói cấu hình rental chi tiết trừ khi được hỏi; chỉ cần nêu các service chính và luồng request.",
    },
    {
        "title": "Slide 12 - Evaluation Setup",
        "goal": "Nói rõ đánh giá là sign-centric 20 samples với LLM judge.",
        "script": (
            "Với phần free-form QA, nhóm đánh giá trên 20 sign-centric samples bằng LLM judge. Các metric được chọn để phản ánh đúng mục tiêu hệ thống: "
            "answer correctness đo câu trả lời có đúng không, citation accuracy đo citation có khớp căn cứ không, evidence faithfulness đo câu trả lời có bám evidence không, "
            "visual grounding proxy đo khả năng bám biển báo, Retrieval Recall@1/3/5 đo khả năng lấy đúng evidence, JSON validity đo output có parse được không, hallucination rate đo lỗi bịa căn cứ, "
            "và average latency đo thời gian phản hồi trung bình."
        ),
        "transition": "Kết quả định lượng chính được tóm tắt ở slide tiếp theo.",
        "avoid": "Không gọi đây là đánh giá toàn diện cho mọi tình huống giao thông; nên gọi là evaluation slice tập trung vào sign-centric QA.",
    },
    {
        "title": "Slide 13 - Final Metrics",
        "goal": "Đọc kết quả nhưng không overclaim.",
        "script": (
            "Kết quả hiện tại cho thấy answer correctness đạt 0.80, citation accuracy đạt 0.65 và evidence faithfulness đạt 1.00. "
            "JSON validity đạt 1.00 và hallucination rate là 0.00 trong tập đánh giá này, nghĩa là output có thể parse ổn định và không ghi nhận trường hợp bịa căn cứ theo judge. "
            "Visual grounding proxy đạt 0.575, còn Retrieval Recall@1, @3 và @5 lần lượt là 0.475, 0.525 và 0.550. Average latency khoảng 17.20 giây. "
            "Nhìn chung, hệ thống thể hiện tốt ở việc sinh câu trả lời có căn cứ, nhưng retrieval và visual grounding vẫn là các phần có nhiều dư địa cải thiện."
        ),
        "transition": "Từ kết quả đó, em kết luận bằng các giới hạn hiện tại và hướng phát triển tiếp theo.",
        "avoid": "Không nói hệ thống đạt accuracy 80% trên toàn bộ dataset; chỉ nói trong tập 20 sign-centric samples đã đánh giá.",
    },
    {
        "title": "Slide 14 - Limitations and Future Work",
        "goal": "Nhận giới hạn một cách chủ động nhưng không tự làm khó.",
        "script": (
            "Hệ thống hiện tại vẫn có một số giới hạn. Thứ nhất, chất lượng phụ thuộc vào detection và retrieval; nếu biển bị crop sai hoặc evidence chưa đúng trọng tâm, answer có thể bị ảnh hưởng. "
            "Thứ hai, một số biển báo có hình thức rất giống nhau nên cần metadata và mô tả ý nghĩa biển đủ rõ. Thứ ba, evaluation hiện tại còn nhỏ và nên được mở rộng với nhiều loại ảnh, ánh sáng và câu hỏi hơn. "
            "Trong tương lai, nhóm sẽ ưu tiên cải thiện sign metadata, mở rộng tập đánh giá, bổ sung kiểm thử thủ công có kiểm soát và tối ưu retrieval để tăng citation accuracy và visual grounding."
        ),
        "transition": "Đó là toàn bộ phần trình bày của nhóm em. Nhóm em xin cảm ơn thầy cô và sẵn sàng trả lời câu hỏi.",
        "avoid": "Không hứa sẽ giải quyết mọi lỗi bằng cách đổi model; nên nói cải thiện retrieval, metadata và evaluation trước.",
    },
]


QA = [
    (
        "Vì sao không chỉ dùng LLM trả lời trực tiếp từ ảnh?",
        "Vì câu hỏi cần căn cứ pháp lý. Nếu chỉ dùng LLM trực tiếp, model có thể trả lời nghe hợp lý nhưng khó kiểm tra. Pipeline hiện tại đưa legal evidence vào prompt và yêu cầu citation theo law_id, article_id để người dùng kiểm chứng.",
    ),
    (
        "Nếu retrieval lấy sai điều luật thì sao?",
        "Đây là một giới hạn đã được nhóm thừa nhận. Vì vậy hệ thống hiển thị citations và evidence để kiểm tra, đồng thời prompt yêu cầu model chỉ dùng context được cung cấp. Hướng cải thiện là tăng chất lượng sign metadata, retrieval ranking và evaluation coverage.",
    ),
    (
        "20 mẫu đánh giá có đủ chưa?",
        "20 mẫu là một evaluation slice tập trung vào sign-centric QA để kiểm tra end-to-end behavior, JSON validity, citation và latency. Nó chưa đại diện cho toàn bộ miền bài toán, nên nhóm trình bày kết quả như một đánh giá có kiểm soát chứ không overclaim.",
    ),
    (
        "Tại sao citation accuracy chưa cao bằng evidence faithfulness?",
        "Evidence faithfulness đo việc câu trả lời có bám vào evidence được cung cấp hay không. Citation accuracy khó hơn vì citation phải khớp đúng article ID và đúng căn cứ. Một câu trả lời có thể bám evidence nhưng citation chưa tối ưu nếu retrieval hoặc mapping article chưa hoàn toàn chính xác.",
    ),
    (
        "Hệ thống có tư vấn pháp lý chính thức được không?",
        "Không. Đây là prototype nghiên cứu và hỗ trợ học tập. Hệ thống giúp người dùng hiểu tình huống và xem căn cứ liên quan, nhưng các quyết định quan trọng vẫn cần đối chiếu văn bản luật gốc hoặc nguồn pháp lý chính thức.",
    ),
    (
        "Vì sao dùng local GPU LLM thay vì API cloud?",
        "Local GPU giúp demo chủ động hơn, giảm phụ thuộc vào external API và giữ pipeline có thể triển khai trong môi trường kiểm soát. Project vẫn dùng OpenAI-compatible interface nên có thể thay endpoint khi cần.",
    ),
    (
        "Điểm mạnh nhất của project là gì?",
        "Điểm mạnh là hệ thống end-to-end: từ ảnh, detect biển, retrieval evidence, resolve điều luật, đến answer JSON có citations và visual evidence. Điều này giúp kết quả dễ kiểm tra hơn so với chỉ trả lời tự nhiên.",
    ),
    (
        "Phần cần cải thiện nhất là gì?",
        "Hai phần đáng cải thiện nhất là retrieval/citation accuracy và visual grounding cho các biển báo dễ nhầm. Đây cũng là lý do nhóm ưu tiên metadata biển báo, sign description và mở rộng evaluation trong future work.",
    ),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D8E2F0"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = BLUE
    r.font.size = Pt(10)
    return p


def add_para(doc, text, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = color or DARK
    return p


def add_callout(doc, label, text, fill="F5F9FF"):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.allow_autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = BLUE
    r.font.size = Pt(10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)
    r2.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_slide(doc, item, index):
    if index > 1:
        doc.add_section(WD_SECTION.NEW_PAGE)
    h = doc.add_heading(item["title"], level=1)
    for run in h.runs:
        run.font.color.rgb = BLUE
        run.font.size = Pt(16)
    add_label(doc, "Mục tiêu slide")
    add_para(doc, item["goal"], MUTED)
    add_label(doc, "Script thuyết trình")
    add_para(doc, item["script"])
    add_label(doc, "Câu chuyển tiếp")
    add_para(doc, item["transition"], MUTED)
    add_callout(doc, "Lưu ý tránh nói hớ", item["avoid"], "FFF8E8")


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_front_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("SCRIPT THUYẾT TRÌNH")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Traffic Sign Legal Question Answering")
    r.font.size = Pt(16)
    r.font.color.rgb = DARK

    add_callout(
        doc,
        "Mục tiêu của script",
        (
            "Bản này được viết để trình bày vừa đủ cho buổi bảo vệ: nói rõ bài toán, pipeline, demo và kết quả; "
            "đồng thời tránh các chi tiết không cần thiết có thể dẫn tới câu hỏi khó ngoài phạm vi hệ thống hiện tại."
        ),
        "F5F9FF",
    )

    add_label(doc, "Cách dùng")
    add_para(
        doc,
        (
            "Mỗi slide gồm bốn phần: mục tiêu slide, lời thoại chính, câu chuyển tiếp và lưu ý tránh nói hớ. "
            "Khi thuyết trình, bạn chỉ cần đọc tự nhiên phần Script thuyết trình và dùng các phần còn lại để luyện tập."
        ),
    )

    add_label(doc, "Nhịp trình bày đề xuất")
    for text in [
        "Tổng thời lượng phù hợp: khoảng 10-12 phút cho phần chính, chưa tính demo và Q&A.",
        "Không mở debug sâu trong demo nếu không được hỏi; chỉ mở citation/evidence để chứng minh câu trả lời có căn cứ.",
        "Khi nói về kết quả, luôn nhắc đây là evaluation slice 20 sign-centric samples, không phải claim cho mọi tình huống giao thông.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(text).font.size = Pt(10.5)


def add_safe_qa(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)
    h = doc.add_heading("Q&A an toàn khi bảo vệ", level=1)
    for run in h.runs:
        run.font.color.rgb = BLUE
        run.font.size = Pt(16)
    add_para(
        doc,
        (
            "Các câu dưới đây giúp trả lời đúng trọng tâm mà không tự mở rộng sang các chi tiết khó bảo vệ. "
            "Bạn có thể dùng nguyên ý, không cần đọc từng chữ."
        ),
        MUTED,
    )
    for question, answer in QA:
        add_label(doc, question)
        add_para(doc, answer)

    add_callout(
        doc,
        "Một câu kết luận nên nhớ",
        (
            "Hệ thống hiện tại mạnh ở khả năng nối ảnh, retrieval, điều luật và citation thành một câu trả lời có thể kiểm tra; "
            "còn phần cần cải thiện chính là retrieval/citation accuracy và visual grounding ở các biển báo dễ nhầm."
        ),
        "F5F9FF",
    )


def build():
    doc = Document()
    configure_styles(doc)
    add_front_page(doc)
    for i, item in enumerate(SLIDES, start=1):
        add_slide(doc, item, i)
    add_safe_qa(doc)
    doc.core_properties.title = "Traffic Sign Legal QA Presentation Script"
    doc.core_properties.subject = "Presentation script for final project defense"
    doc.core_properties.author = "Traffic Sign Legal QA Team"
    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_REPO)
    doc.save(OUT_DOWNLOADS)
    print(OUT_REPO)
    print(OUT_DOWNLOADS)


if __name__ == "__main__":
    build()
